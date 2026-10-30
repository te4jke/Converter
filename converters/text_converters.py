import html
from html.parser import HTMLParser

from docx import Document

from converters.common import read_text, write_text
from core.exceptions import ConversionFailedError

HTML_TEMPLATE = (
    "<!DOCTYPE html>\n<html>\n<head>\n<meta charset=\"utf-8\">\n"
    "<title>{title}</title>\n</head>\n<body>\n{body}\n</body>\n</html>\n"
)

BLOCK_TAGS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5",
              "h6", "ul", "ol", "table", "section", "article", "pre"}
SKIP_TAGS = {"script", "style", "head"}


def txt_to_docx(src, dst):
    text = read_text(src)
    try:
        doc = Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(dst)
    except ValueError:
        raise ConversionFailedError("В тексте есть недопустимые символы")
    except OSError as e:
        raise ConversionFailedError("Не удалось сохранить документ: " + str(e))


def docx_to_txt(src, dst):
    try:
        doc = Document(src)
    except Exception as e:
        raise ConversionFailedError("Не удалось открыть документ: " + str(e))

    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))

    write_text(dst, "\n".join(lines))


def txt_to_html(src, dst):
    text = read_text(src)
    paragraphs = []
    for line in text.splitlines():
        paragraphs.append("<p>" + html.escape(line) + "</p>")

    page = HTML_TEMPLATE.format(title="", body="\n".join(paragraphs))
    write_text(dst, page)


class _TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in SKIP_TAGS:
            self.skip += 1
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in SKIP_TAGS:
            self.skip = max(0, self.skip - 1)
        elif tag in BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def html_to_txt(src, dst):
    parser = _TextExtractor()
    parser.feed(read_text(src))
    parser.close()

    lines = [" ".join(line.split()) for line in "".join(parser.parts).split("\n")]
    write_text(dst, "\n".join(line for line in lines if line) + "\n")
