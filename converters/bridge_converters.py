# Мосты между группами форматов: текст и документы <-> таблицы и данные.
# Благодаря им цепочки могут проходить, например, TXT -> JSON -> XLSX

import csv
import html
import io
import json
from html.parser import HTMLParser

import pandas as pd
from docx import Document

from converters.common import read_text, write_text
from converters.json_converters import _cell
from converters.text_converters import HTML_TEMPLATE
from core.exceptions import ConversionFailedError


def _read_csv_rows(path):
    text = read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    try:
        return list(csv.reader(io.StringIO(text), dialect=dialect))
    except csv.Error as e:
        raise ConversionFailedError("Некорректный CSV: " + str(e))


def _read_json_rows(path):
    try:
        data = json.loads(read_text(path))
    except json.JSONDecodeError as e:
        raise ConversionFailedError("Некорректный JSON: " + str(e))
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise ConversionFailedError("Для таблицы нужен JSON-объект или список объектов")
    fields = []
    for row in data:
        for key in row:
            if key not in fields:
                fields.append(key)
    return [fields] + [[_cell(r.get(k)) for k in fields] for r in data]


def _read_xlsx_rows(path):
    try:
        df = pd.read_excel(path, engine="openpyxl", header=None)
    except Exception as e:
        raise ConversionFailedError("Не удалось прочитать таблицу: " + str(e))
    df = df.fillna("")
    return [[str(c) for c in row] for row in df.values.tolist()]


def _rows_to_txt(rows):
    return "\n".join("\t".join(str(c) for c in row) for row in rows) + "\n"


def _rows_to_html(rows):
    if not rows:
        raise ConversionFailedError("В таблице нет данных")
    out = ["<table border=\"1\" cellspacing=\"0\" cellpadding=\"4\">"]
    for i, row in enumerate(rows):
        tag = "th" if i == 0 else "td"
        cells = "".join("<%s>%s</%s>" % (tag, html.escape(str(c)), tag) for c in row)
        out.append("<tr>%s</tr>" % cells)
    out.append("</table>")
    return HTML_TEMPLATE.format(title="", body="\n".join(out))


def _write_rows_csv(rows, dst):
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(rows)
    write_text(dst, out.getvalue())


def txt_to_json(src, dst):
    # Каждая строка файла становится элементом массива
    lines = read_text(src).splitlines()
    if not "".join(lines).strip():
        raise ConversionFailedError("Файл пустой")
    write_text(dst, json.dumps(lines, ensure_ascii=False, indent=2) + "\n")


def json_to_txt(src, dst):
    try:
        data = json.loads(read_text(src))
    except json.JSONDecodeError as e:
        raise ConversionFailedError("Некорректный JSON: " + str(e))
    # Массив строк — обратное преобразование к txt_to_json
    if isinstance(data, list) and all(isinstance(x, str) for x in data):
        write_text(dst, "\n".join(data) + "\n")
    else:
        write_text(dst, json.dumps(data, ensure_ascii=False, indent=2) + "\n")


def txt_to_csv(src, dst):
    # Строка файла = строка таблицы, разделитель определяется автоматически
    text = read_text(src)
    if not text.strip():
        raise ConversionFailedError("Файл пустой")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = None
    lines = text.splitlines()
    if dialect:
        rows = list(csv.reader(lines, dialect=dialect))
    else:
        rows = [[line] for line in lines]
    _write_rows_csv(rows, dst)


def csv_to_txt(src, dst):
    write_text(dst, _rows_to_txt(_read_csv_rows(src)))


def xlsx_to_txt(src, dst):
    write_text(dst, _rows_to_txt(_read_xlsx_rows(src)))


def csv_to_html(src, dst):
    write_text(dst, _rows_to_html(_read_csv_rows(src)))


def json_to_html(src, dst):
    write_text(dst, _rows_to_html(_read_json_rows(src)))


def xlsx_to_html(src, dst):
    write_text(dst, _rows_to_html(_read_xlsx_rows(src)))


def docx_to_csv(src, dst):
    # Берётся первая таблица документа
    try:
        doc = Document(src)
    except Exception as e:
        raise ConversionFailedError("Не удалось открыть документ: " + str(e))
    if not doc.tables:
        raise ConversionFailedError("В документе нет таблиц")
    rows = [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    _write_rows_csv(rows, dst)


class _TableExtractor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows, self.row, self.cell = [], None, None
        self.depth = 0
        self.done = False

    def handle_starttag(self, tag, attrs):
        if self.done:
            return
        if tag == "table":
            self.depth += 1
        elif self.depth == 1 and tag == "tr":
            self.row = []
        elif self.depth == 1 and tag in ("td", "th") and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if self.done:
            return
        if tag in ("td", "th") and self.depth == 1 and self.cell is not None and self.row is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif tag == "tr" and self.depth == 1 and self.row is not None:
            self.rows.append(self.row)
            self.row = None
        elif tag == "table":
            self.depth -= 1
            if self.depth == 0:
                self.done = True

    def handle_data(self, data):
        # текст вложенных таблиц в ячейку внешней не попадает
        if self.cell is not None and self.depth == 1:
            self.cell.append(data)


def html_to_csv(src, dst):
    # Берётся первая таблица страницы
    parser = _TableExtractor()
    parser.feed(read_text(src))
    parser.close()
    if not parser.rows:
        raise ConversionFailedError("На странице нет таблиц")
    _write_rows_csv(parser.rows, dst)
