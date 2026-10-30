import io
import os
import shutil
import tempfile
import unittest

from docx import Document
from PIL import Image
from pypdf import PdfReader


class TempDirCase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)

    def path(self, name):
        return os.path.join(self.dir, name)

    def write(self, name, content):
        mode = "wb" if isinstance(content, bytes) else "w"
        kwargs = {} if mode == "wb" else {"encoding": "utf-8"}
        with open(self.path(name), mode, **kwargs) as f:
            f.write(content)
        return self.path(name)

    def read(self, name):
        with open(self.path(name), encoding="utf-8") as f:
            return f.read()


def png_bytes(color="red", size=(40, 30)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def docx_bytes(*paragraphs, table=None):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, value in enumerate(row):
                t.cell(i, j).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pdf_pages(path):
    return len(PdfReader(path).pages)
